import os
import subprocess
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_lab_manual(output_docx_path):
    doc = docx.Document()

    # Page Margins
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    # Styles
    # Title Header
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = title_p.add_run("SAVEETHA INSTITUTE OF MEDICAL AND TECHNICAL SCIENCES\n")
    r1.bold = True
    r1.font.size = Pt(14)
    r1.font.name = "Calibri"
    r1.font.color.rgb = RGBColor(0, 51, 102)

    r2 = title_p.add_run("SIMATS ENGINEERING\nDEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING\n")
    r2.bold = True
    r2.font.size = Pt(12)
    r2.font.name = "Calibri"
    r2.font.color.rgb = RGBColor(0, 102, 153)

    r3 = title_p.add_run("CSA17 — ARTIFICIAL INTELLIGENCE LAB OBSERVATION MANUAL\n")
    r3.bold = True
    r3.font.size = Pt(13)
    r3.font.name = "Calibri"
    r3.font.color.rgb = RGBColor(180, 50, 20)

    # Student Info Table
    info_table = doc.add_table(rows=2, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_table.autofit = False

    data = [
        ("STUDENT NAME: VISHAL BALAJI J", "REGISTER NUMBER: 192525014"),
        ("COURSE CODE: CSA17 - AI", "DEGREE & BRANCH: B.TECH CSE")
    ]
    for row_idx, (c1_text, c2_text) in enumerate(data):
        row = info_table.rows[row_idx]
        cell1, cell2 = row.cells[0], row.cells[1]
        cell1.width = Inches(3.4)
        cell2.width = Inches(3.4)
        
        p1 = cell1.paragraphs[0]
        r = p1.add_run(c1_text)
        r.bold = True
        r.font.size = Pt(10)
        
        p2 = cell2.paragraphs[0]
        r = p2.add_run(c2_text)
        r.bold = True
        r.font.size = Pt(10)

        set_cell_background(cell1, "F0F4F8")
        set_cell_background(cell2, "F0F4F8")
        set_cell_margins(cell1, top=120, bottom=120, left=150, right=150)
        set_cell_margins(cell2, top=120, bottom=120, left=150, right=150)

    doc.add_paragraph() # Spacing

    # Index Table
    doc.add_heading("INDEX / EXPERIMENT SCHEDULE", level=1)
    
    idx_table = doc.add_table(rows=1, cols=6)
    idx_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Exp. No.", "Date", "Experiment Title", "Course Outcome", "Status", "Sign"]
    hdr_widths = [Inches(0.8), Inches(1.0), Inches(2.6), Inches(0.8), Inches(0.9), Inches(0.7)]
    
    hdr_row = idx_table.rows[0]
    for i, h_text in enumerate(headers):
        cell = hdr_row.cells[i]
        cell.width = hdr_widths[i]
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h_text)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, "003366")
        set_cell_margins(cell, top=100, bottom=100, left=80, right=80)

    index_data = [
        ("1", "13.10.2026", "8-Puzzle Problem using A* Search", "CO1", "Completed", ""),
        ("2", "13.10.2026", "8-Queen Problem using Backtracking", "CO1", "Completed", ""),
        ("3", "13.10.2026", "Water Jug Problem using State Space Search", "CO2", "Completed", ""),
        ("4", "13.10.2026", "Missionaries and Cannibals Problem", "CO3", "Completed", ""),
        ("5", "13.10.2026", "Vacuum Cleaner Agent (Vacuum World)", "CO3", "Completed", ""),
        ("6", "15.10.2026", "Breadth First Search (BFS) Implementation", "CO3", "Completed", ""),
        ("7", "15.10.2026", "Depth First Search (DFS) Implementation", "CO4", "Completed", ""),
        ("8", "15.10.2026", "A* Search Algorithm on Weighted Graph", "CO4", "Completed", ""),
        ("9", "15.10.2026", "Crypt-Arithmetic Problem (SEND+MORE=MONEY)", "CO2", "Completed", ""),
        ("10", "15.10.2026", "Travelling Salesman Problem (TSP)", "CO4", "Completed", ""),
    ]

    for row_idx, row_vals in enumerate(index_data):
        row = idx_table.add_row()
        bg_col = "FFFFFF" if row_idx % 2 == 0 else "F7FAFC"
        for i, val in enumerate(row_vals):
            cell = row.cells[i]
            cell.width = hdr_widths[i]
            p = cell.paragraphs[0]
            if i in [0, 1, 3, 4]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(val)
            r.font.size = Pt(9)
            set_cell_background(cell, bg_col)
            set_cell_margins(cell, top=80, bottom=80, left=80, right=80)

    doc.add_page_break()

    # Directory of source code
    base_dir = r"c:\Users\Vishal_balaji_J\Downloads\git hub\Artificial-Intelligence\Lab Experiments"

    # Experiments metadata
    experiments = [
        {
            "num": 1,
            "title": "8-PUZZLE PROBLEM USING A* SEARCH",
            "date": "13.10.2026",
            "co": "CO1",
            "aim": "To write and execute a Python program to solve the 8-puzzle problem using the A* Search Algorithm with the Manhattan Distance heuristic.",
            "procedure": [
                "Understand the 8-puzzle board configuration consisting of a 3x3 grid with tiles numbered 1 to 8 and one empty blank space (represented by 0).",
                "Define the initial state and the target goal state.",
                "Implement the Manhattan Distance heuristic function h(n) calculating the sum of absolute horizontal and vertical distances of tiles from their goal positions.",
                "Generate valid successor states by sliding the blank tile Up, Down, Left, or Right.",
                "Employ a priority queue (min-heap) to explore states with the minimum evaluation cost f(n) = g(n) + h(n).",
                "Maintain a visited set to avoid redundant state expansions and cycles.",
                "Execute the program in Python IDLE and record the optimal sequence of puzzle configurations."
            ],
            "algorithm": [
                "Step 1: Start the program.",
                "Step 2: Define the Goal State as [[1, 2, 3], [4, 5, 6], [7, 8, 0]].",
                "Step 3: Define heuristic function h(state) returning the total Manhattan distance of all tiles from their goal coordinates.",
                "Step 4: Initialize a Priority Queue (Open List) and push the Initial State with g = 0 and f = h(initial).",
                "Step 5: Initialize a Visited Dictionary (Closed List) to track minimum path costs.",
                "Step 6: While the Priority Queue is not empty:",
                "       a. Pop the node with the minimum f-score.",
                "       b. If current state equals Goal State, reconstruct the solution path and terminate.",
                "       c. For each valid move (UP, DOWN, LEFT, RIGHT):",
                "          i. Calculate tentative cost g_new = g + 1.",
                "          ii. If neighbor is unvisited or g_new < visited[neighbor], update visited and push to Priority Queue.",
                "Step 7: Display the step-by-step tile moves and final solved board.",
                "Step 8: Stop the program."
            ],
            "file": "Exp1_8_Puzzle.py",
            "result": "The Python program to solve the 8-puzzle problem using the A* Search algorithm was successfully executed in Python IDLE. The optimal path to reach the goal state was determined in 3 steps."
        },
        {
            "num": 2,
            "title": "8-QUEEN PROBLEM USING BACKTRACKING",
            "date": "13.10.2026",
            "co": "CO1",
            "aim": "To write and execute a Python program to solve the 8-Queen problem using the Backtracking algorithm.",
            "procedure": [
                "Understand the problem: Place 8 queens on an 8x8 chessboard such that no two queens attack each other.",
                "Formulate non-attacking constraints: No two queens can share the same row, same column, or same diagonal.",
                "Develop the is_safe() utility function to check whether a queen can be safely placed in board[row][col] without conflicting with previously placed queens.",
                "Use a recursive backtracking approach to place queens column by column.",
                "If placing a queen leads to a valid state, proceed to the next column. If no safe row exists, backtrack by removing the previously placed queen.",
                "Execute the code in Python IDLE, find all valid solutions (92 configurations), and display the board matrices.",
                "Record the output and verify the queen coordinates."
            ],
            "algorithm": [
                "Step 1: Start the program.",
                "Step 2: Initialize an 8x8 chessboard matrix with all cells set to 0.",
                "Step 3: Define function is_safe(board, row, col):",
                "       a. Check if any queen exists in the same row on the left columns.",
                "       b. Check if any queen exists on the upper diagonal to the left.",
                "       c. Check if any queen exists on the lower diagonal to the left.",
                "       d. Return True if all checks pass; otherwise return False.",
                "Step 4: Define recursive function backtrack(col):",
                "       a. If col >= 8, all 8 queens are safely placed; save the board configuration.",
                "       b. For each row from 0 to 7:",
                "          i. If is_safe(board, row, col) is True:",
                "             - Place queen: board[row][col] = 1.",
                "             - Recursively call backtrack(col + 1).",
                "             - Backtrack: board[row][col] = 0.",
                "Step 5: Display the visual 8x8 board representation and queen coordinates.",
                "Step 6: Stop the program."
            ],
            "file": "Exp2_8_Queen.py",
            "result": "The Python program for the 8-Queen problem was successfully developed and executed in Python IDLE. All 92 non-attacking queen solutions were identified and displayed."
        },
        {
            "num": 3,
            "title": "WATER JUG PROBLEM USING STATE SPACE SEARCH",
            "date": "13.10.2026",
            "co": "CO2",
            "aim": "To write and execute a Python program to solve the Water Jug Problem using state space search (Breadth First Search).",
            "procedure": [
                "Identify the given parameters: Jug A capacity = 4 Liters, Jug B capacity = 3 Liters, Target quantity = 2 Liters in Jug A.",
                "Model state as a tuple (a, b) representing the current volume of water in Jug A and Jug B.",
                "Define production rules: Fill Jug A, Fill Jug B, Empty Jug A, Empty Jug B, Pour Jug A to Jug B, Pour Jug B to Jug A.",
                "Use Breadth-First Search (BFS) with a FIFO queue to systematically explore reachable states starting from (0, 0).",
                "Check for the goal state (2, b) and reconstruct the step-by-step operations.",
                "Run the program in Python IDLE and observe the state transitions."
            ],
            "algorithm": [
                "Step 1: Start the program.",
                "Step 2: Initialize capacities cap_a = 4, cap_b = 3, target = 2.",
                "Step 3: Initialize BFS Queue with initial state ((0, 0), []) and a visited set.",
                "Step 4: While Queue is not empty:",
                "       a. Dequeue current state (a, b) and operation path.",
                "       b. If a == target, return the solution path.",
                "       c. Generate all 6 possible successor states:",
                "          1. Fill A: (cap_a, b)",
                "          2. Fill B: (a, cap_b)",
                "          3. Empty A: (0, b)",
                "          4. Empty B: (a, 0)",
                "          5. Pour A -> B: (a - pour, b + pour) where pour = min(a, cap_b - b)",
                "          6. Pour B -> A: (a + pour, b - pour) where pour = min(b, cap_a - a)",
                "       d. For each unvisited successor state, add to visited set and enqueue.",
                "Step 5: Print the formatted table of steps, actions taken, and jug volumes.",
                "Step 6: Stop the program."
            ],
            "file": "Exp3_Water_Jug.py",
            "result": "The Water Jug problem was successfully solved in Python using BFS. Exactly 2 Liters were measured in Jug A in 6 operational steps without measuring marks."
        },
        {
            "num": 4,
            "title": "MISSIONARIES AND CANNIBALS PROBLEM",
            "date": "13.10.2026",
            "co": "CO3",
            "aim": "To write and execute a Python program to solve the Missionaries and Cannibals problem using state space search.",
            "procedure": [
                "Understand the problem constraints: 3 Missionaries (M) and 3 Cannibals (C) are on the left bank and must cross to the right bank using a boat carrying at most 2 people.",
                "Safety rule: If missionaries are present on either bank, they must never be outnumbered by cannibals (M >= C).",
                "Represent state as (m_left, c_left, boat_pos) where boat_pos is 1 (Left bank) or 0 (Right bank).",
                "Define valid boat movements: (1,0), (2,0), (0,1), (0,2), (1,1).",
                "Implement BFS to find the shortest safe crossing sequence from (3,3,1) to (0,0,0).",
                "Execute the program in Python IDLE and record the river crossing schedule."
            ],
            "algorithm": [
                "Step 1: Start the program.",
                "Step 2: Define initial state as (3, 3, 1) and goal state as (0, 0, 0).",
                "Step 3: Define validation function is_valid(m, c):",
                "       - If m > 0 and m < c: return False (Missionaries eaten on left bank).",
                "       - If (3 - m) > 0 and (3 - m) < (3 - c): return False (Missionaries eaten on right bank).",
                "       - Else return True.",
                "Step 4: Initialize BFS queue with (initial_state, []) and add initial state to visited set.",
                "Step 5: While queue is not empty:",
                "       a. Dequeue (current_state, path).",
                "       b. If current_state == (0, 0, 0), return path.",
                "       c. For each possible boat move (dm, dc) in [(1,0), (2,0), (0,1), (0,2), (1,1)]:",
                "          - Compute next state depending on boat direction.",
                "          - If valid and unvisited, add to visited set and enqueue.",
                "Step 6: Display the step-by-step river crossing log.",
                "Step 7: Stop the program."
            ],
            "file": "Exp4_Missionaries_Cannibals.py",
            "result": "The Missionaries and Cannibals problem was successfully solved using Breadth-First Search in Python IDLE. All 6 individuals crossed safely in 11 boat trips."
        },
        {
            "num": 5,
            "title": "VACUUM CLEANER PROBLEM (VACUUM WORLD)",
            "date": "13.10.2026",
            "co": "CO3",
            "aim": "To write and execute a Python program to implement the Vacuum Cleaner reflex agent in a two-room environment.",
            "procedure": [
                "Formulate the environment: Two rooms A and B, where each room can be either 'Clean' or 'Dirty'.",
                "Define the agent percept: Current location (A or B) and current room status ('Clean' or 'Dirty').",
                "Formulate the condition-action reflex rules:",
                "   - If current room is 'Dirty', perform action 'SUCK'.",
                "   - If current room is 'Clean' and location is 'A', perform action 'RIGHT'.",
                "   - If current room is 'Clean' and location is 'B', perform action 'LEFT'.",
                "Simulate agent interactions under multiple starting conditions (both dirty, one clean, etc.).",
                "Record the action logs and environment status changes until all rooms are clean."
            ],
            "algorithm": [
                "Step 1: Start the program.",
                "Step 2: Create VacuumCleanerAgent class with location and cost tracker.",
                "Step 3: Define reflex_action(status) method:",
                "       - If status == 'Dirty', return 'SUCK'.",
                "       - Else if location == 'A', return 'RIGHT'.",
                "       - Else if location == 'B', return 'LEFT'.",
                "Step 4: Define simulation loop:",
                "       - While Room A is 'Dirty' or Room B is 'Dirty':",
                "         a. Sense current room status.",
                "         b. Invoke reflex_action().",
                "         c. Update room status to 'Clean' if 'SUCK' was chosen.",
                "         d. Update agent position if movement action was chosen.",
                "         e. Increment agent cost counter and log state.",
                "Step 5: Print final summary showing all rooms clean and total cost.",
                "Step 6: Stop the program."
            ],
            "file": "Exp5_Vacuum_Cleaner.py",
            "result": "The Vacuum Cleaner reflex agent was successfully simulated in Python IDLE. The agent effectively perceived its environment, cleaned both rooms, and minimized movement costs across multiple test scenarios."
        },
        {
            "num": 6,
            "title": "BREADTH FIRST SEARCH (BFS) IMPLEMENTATION",
            "date": "15.10.2026",
            "co": "CO3",
            "aim": "To write and execute a Python program to implement the Breadth First Search (BFS) graph traversal algorithm.",
            "procedure": [
                "Represent the graph using an Adjacency List (Python dictionary).",
                "Select a starting node to begin traversal.",
                "Use a FIFO (First-In, First-Out) queue from the collections module (deque) to manage frontier nodes.",
                "Maintain a visited list to prevent revisiting already explored nodes.",
                "Explore all neighbor nodes level by level, printing queue states at each step.",
                "Implement a shortest path utility using BFS.",
                "Execute the program in Python IDLE and record the traversal sequence."
            ],
            "algorithm": [
                "Step 1: Start the program.",
                "Step 2: Represent graph as adjacency list with nodes and neighbor lists.",
                "Step 3: Define bfs_traversal(graph, start_node):",
                "       a. Initialize empty list visited = [].",
                "       b. Initialize FIFO queue with start_node.",
                "       c. Mark start_node as seen.",
                "Step 4: While queue is not empty:",
                "       a. Dequeue current node.",
                "       b. Append current node to visited list.",
                "       c. For each unvisited neighbor of current node:",
                "          - Mark neighbor as seen.",
                "          - Enqueue neighbor into queue.",
                "Step 5: Output the step-by-step queue contents and final traversal path.",
                "Step 6: Stop the program."
            ],
            "file": "Exp6_BFS.py",
            "result": "The Breadth First Search (BFS) algorithm was successfully implemented and executed in Python IDLE. The graph was traversed level by level producing the order: A -> B -> C -> D -> E -> F."
        },
        {
            "num": 7,
            "title": "DEPTH FIRST SEARCH (DFS) IMPLEMENTATION",
            "date": "15.10.2026",
            "co": "CO4",
            "aim": "To write and execute a Python program to implement the Depth First Search (DFS) graph traversal algorithm.",
            "procedure": [
                "Represent the directed/undirected graph as an adjacency list.",
                "Implement DFS using both recursive call stack and explicit LIFO stack.",
                "Begin exploration from a designated root node.",
                "Traverse as deeply as possible along each branch before backtracking.",
                "Maintain a visited set to avoid infinite loops and redundant exploration.",
                "Run the program in Python IDLE, compare iterative and recursive results, and record outputs."
            ],
            "algorithm": [
                "Step 1: Start the program.",
                "Step 2: Define graph structure using an adjacency dictionary.",
                "Step 3: Recursive Approach:",
                "       a. Define dfs_recursive(graph, node, visited):",
                "       b. Mark node as visited.",
                "       c. For each unvisited neighbor, recursively call dfs_recursive().",
                "Step 4: Iterative Approach (LIFO Stack):",
                "       a. Initialize stack = [start_node] and seen = set().",
                "       b. While stack is not empty:",
                "          i. Pop top node from stack.",
                "          ii. If not in seen, mark as seen and add to visited list.",
                "          iii. Push unvisited neighbors onto stack in reverse order.",
                "Step 5: Display the step-by-step stack states and final traversal order.",
                "Step 6: Stop the program."
            ],
            "file": "Exp7_DFS.py",
            "result": "The Depth First Search (DFS) algorithm was successfully implemented in Python IDLE using both recursive and iterative stack methods. Both methods verified identical deep traversal paths."
        },
        {
            "num": 8,
            "title": "A* SEARCH ALGORITHM IMPLEMENTATION",
            "date": "15.10.2026",
            "co": "CO4",
            "aim": "To write and execute a Python program to implement the A* Search Algorithm on a weighted graph with heuristics.",
            "procedure": [
                "Construct a weighted graph with edge costs g(n) and heuristic estimates h(n) to the goal node.",
                "Understand the evaluation function f(n) = g(n) + h(n), balancing actual path cost and heuristic prediction.",
                "Initialize a Priority Queue (Open List) containing the start node.",
                "Maintain a Closed List to store already evaluated nodes and prevent suboptimal re-expansions.",
                "Pop the node with minimum f(n), examine outgoing edges, and update path costs.",
                "Terminate when the goal node is popped from the priority queue.",
                "Display expansion steps, optimal path, and minimum total cost."
            ],
            "algorithm": [
                "Step 1: Start the program.",
                "Step 2: Define graph with weighted edges and heuristic table h(n).",
                "Step 3: Define a_star_search(graph, heuristics, start, goal):",
                "       a. Push start node to priority queue with f = h(start), g = 0.",
                "       b. Initialize g_scores dictionary with g_scores[start] = 0.",
                "       c. Initialize closed_set = set().",
                "Step 4: While priority queue is not empty:",
                "       a. Pop node with lowest f-score.",
                "       b. If current node == goal, return reconstructed path and cost g.",
                "       c. Add current node to closed_set.",
                "       d. For each (neighbor, weight) from current node:",
                "          i. Calculate tentative_g = g + weight.",
                "          ii. If tentative_g < g_scores[neighbor]:",
                "              - Update g_scores[neighbor] = tentative_g.",
                "              - Calculate new_f = tentative_g + h(neighbor).",
                "              - Push neighbor into priority queue.",
                "Step 5: Display step-by-step evaluation table and optimal path.",
                "Step 6: Stop the program."
            ],
            "file": "Exp8_A_Star.py",
            "result": "The A* Search algorithm was successfully executed in Python IDLE. The optimal path from Start 'S' to Goal 'G' was determined as S -> A -> B -> C -> D -> G with an optimal total cost of 9."
        },
        {
            "num": 9,
            "title": "CRYPT-ARITHMETIC PROBLEM (SEND + MORE = MONEY)",
            "date": "15.10.2026",
            "co": "CO2",
            "aim": "To write and execute a Python program to solve the Crypt-Arithmetic Problem using Constraint Satisfaction and Backtracking.",
            "procedure": [
                "Understand the puzzle: SEND + MORE = MONEY where each distinct letter represents a unique decimal digit (0 to 9).",
                "Identify constraints:",
                "   - Unique digit per letter across {S, E, N, D, M, O, R, Y}.",
                "   - Leading letters S and M must not be 0.",
                "   - The arithmetic summation constraint: SEND + MORE == MONEY.",
                "Formulate search space using permutations of digits and evaluate validity.",
                "Run the program in Python IDLE and verify the resulting digit assignments."
            ],
            "algorithm": [
                "Step 1: Start the program.",
                "Step 2: Define unique letters list ['S', 'E', 'N', 'D', 'M', 'O', 'R', 'Y'].",
                "Step 3: Iterate through permutations of digits (0 to 9) of length 8.",
                "Step 4: Map each permutation to the letter list.",
                "Step 5: Check leading digit constraint: if S == 0 or M == 0, continue.",
                "Step 6: Calculate numerical values:",
                "       - SEND = 1000*S + 100*E + 10*N + D",
                "       - MORE = 1000*M + 100*O + 10*R + E",
                "       - MONEY = 10000*M + 1000*O + 100*N + 10*E + Y",
                "Step 7: If SEND + MORE == MONEY, output the solution and terminate.",
                "Step 8: Stop the program."
            ],
            "file": "Exp9_Crypt_Arithmetic.py",
            "result": "The Crypt-Arithmetic problem SEND + MORE = MONEY was successfully solved in Python IDLE. The unique assignment (9567 + 1085 = 10652) satisfies all arithmetic and CSP constraints."
        },
        {
            "num": 10,
            "title": "TRAVELLING SALESMAN PROBLEM (TSP)",
            "date": "15.10.2026",
            "co": "CO4",
            "aim": "To write and execute a Python program to solve the Travelling Salesman Problem using State Space Search / Branch and Bound.",
            "procedure": [
                "Understand the TSP problem: Given a set of cities and pairwise distances, find the shortest closed tour visiting each city exactly once and returning to the origin.",
                "Represent city inter-distances using a 4x4 symmetric distance matrix.",
                "Fix the starting city (City 'A') to avoid circular permutation duplicates.",
                "Evaluate all possible permutations of visiting intermediate cities.",
                "Compute total round-trip costs and determine the minimum cost Hamiltonian cycle.",
                "Execute the program in Python IDLE and display all evaluated tours."
            ],
            "algorithm": [
                "Step 1: Start the program.",
                "Step 2: Define list of cities ['A', 'B', 'C', 'D'] and distance matrix.",
                "Step 3: Set start_city = 'A'.",
                "Step 4: Generate all permutations of remaining cities ['B', 'C', 'D'].",
                "Step 5: For each permutation:",
                "       a. Construct tour = [start] + list(perm) + [start].",
                "       b. Calculate tour cost = sum of adjacent edge weights in distance matrix.",
                "       c. If tour cost < min_cost, update min_cost and best_tour.",
                "Step 6: Print table of all candidate tours and highlight the optimal tour.",
                "Step 7: Stop the program."
            ],
            "file": "Exp10_Travelling_Salesman.py",
            "result": "The Travelling Salesman Problem was successfully implemented and solved in Python IDLE. The optimal tour A -> B -> D -> C -> A was identified with a minimum distance of 80."
        }
    ]

    # Process each experiment
    for exp in experiments:
        # Load source code
        py_path = os.path.join(base_dir, exp["file"])
        with open(py_path, 'r', encoding='utf-8') as f:
            code_text = f.read()

        # Run script to get exact IDLE output
        res = subprocess.run(["python", py_path], capture_output=True, text=True, cwd=base_dir)
        idle_output = res.stdout.strip()

        # Add experiment header
        h1 = doc.add_heading(f"EXPERIMENT {exp['num']}: {exp['title']}", level=1)
        h1.paragraph_format.space_before = Pt(14)
        h1.paragraph_format.space_after = Pt(4)

        # Meta table: Date & CO
        meta_table = doc.add_table(rows=1, cols=3)
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        meta_row = meta_table.rows[0]
        meta_row.cells[0].width = Inches(2.2)
        meta_row.cells[1].width = Inches(2.2)
        meta_row.cells[2].width = Inches(2.4)
        
        m_p0 = meta_row.cells[0].paragraphs[0]
        mr0 = m_p0.add_run(f"DATE: {exp['date']}")
        mr0.bold = True
        mr0.font.size = Pt(9.5)
        
        m_p1 = meta_row.cells[1].paragraphs[0]
        mr1 = m_p1.add_run(f"COURSE OUTCOME: {exp['co']}")
        mr1.bold = True
        mr1.font.size = Pt(9.5)

        m_p2 = meta_row.cells[2].paragraphs[0]
        deadline_note = "Assigned for 13.10.26" if "13.10" in exp['date'] else "Assigned for 15.10.26"
        mr2 = m_p2.add_run(f"SCHEDULE: {deadline_note}")
        mr2.bold = True
        mr2.font.size = Pt(9.5)

        for cell in meta_row.cells:
            set_cell_background(cell, "EAEFF5")
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)

        doc.add_paragraph()

        # AIM
        h_aim = doc.add_heading("AIM", level=2)
        h_aim.paragraph_format.space_before = Pt(8)
        p_aim = doc.add_paragraph(exp["aim"])
        p_aim.paragraph_format.line_spacing = 1.15

        # PROCEDURE
        h_proc = doc.add_heading("PROCEDURE", level=2)
        h_proc.paragraph_format.space_before = Pt(8)
        for step in exp["procedure"]:
            p = doc.add_paragraph(step, style='List Bullet')
            p.paragraph_format.space_after = Pt(2)

        # ALGORITHM
        h_algo = doc.add_heading("ALGORITHM", level=2)
        h_algo.paragraph_format.space_before = Pt(8)
        for step in exp["algorithm"]:
            p = doc.add_paragraph(step)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Inches(0.2)

        # SOURCE CODE
        h_code = doc.add_heading("SOURCE CODE (PYTHON)", level=2)
        h_code.paragraph_format.space_before = Pt(8)

        # Boxed table for source code
        code_table = doc.add_table(rows=1, cols=1)
        code_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        code_cell = code_table.rows[0].cells[0]
        code_cell.width = Inches(6.8)
        set_cell_background(code_cell, "F8F9FA")
        set_cell_margins(code_cell, top=100, bottom=100, left=150, right=150)
        
        cp = code_cell.paragraphs[0]
        cp.paragraph_format.line_spacing = 1.0
        cp.paragraph_format.space_after = Pt(0)
        c_run = cp.add_run(code_text)
        c_run.font.name = "Consolas"
        c_run.font.size = Pt(8.5)
        c_run.font.color.rgb = RGBColor(30, 30, 30)

        doc.add_paragraph()

        # PYTHON COMPILER IDLE OUTPUT
        h_out = doc.add_heading("PYTHON COMPILER IDLE OUTPUT", level=2)
        h_out.paragraph_format.space_before = Pt(8)

        # Boxed table for IDLE output (Dark / Terminal theme)
        out_table = doc.add_table(rows=1, cols=1)
        out_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        out_cell = out_table.rows[0].cells[0]
        out_cell.width = Inches(6.8)
        set_cell_background(out_cell, "1E1E1E")
        set_cell_margins(out_cell, top=120, bottom=120, left=150, right=150)

        op = out_cell.paragraphs[0]
        op.paragraph_format.line_spacing = 1.05
        op.paragraph_format.space_after = Pt(0)
        
        # IDLE header prefix
        idle_header = "=================== RESTART: Shell Execution ===================\n"
        o_run = op.add_run(idle_header + idle_output)
        o_run.font.name = "Consolas"
        o_run.font.size = Pt(8.5)
        o_run.font.color.rgb = RGBColor(0, 255, 128) # Classic IDLE / Terminal Green

        doc.add_paragraph()

        # RESULT
        h_res = doc.add_heading("RESULT", level=2)
        h_res.paragraph_format.space_before = Pt(8)
        p_res = doc.add_paragraph(exp["result"])
        p_res.paragraph_format.line_spacing = 1.15
        p_res.paragraph_format.space_after = Pt(14)

        # Page break between experiments
        if exp["num"] < len(experiments):
            doc.add_page_break()

    # Save Document
    doc.save(output_docx_path)
    print(f"Successfully generated Lab Manual at: {output_docx_path}")

if __name__ == '__main__':
    downloads_path = r"c:\Users\Vishal_balaji_J\Downloads\CSA17_Artificial_Intelligence_Lab_Observation_Manual.docx"
    repo_path = r"c:\Users\Vishal_balaji_J\Downloads\git hub\Artificial-Intelligence\Lab Experiments\CSA17_Artificial_Intelligence_Lab_Observation_Manual.docx"
    repo2_path = r"c:\Users\Vishal_balaji_J\Downloads\git hub\Artificial-intelligence_CSA-1709\Lab Experiments\CSA17_Artificial_Intelligence_Lab_Observation_Manual.docx"

    create_lab_manual(downloads_path)
    create_lab_manual(repo_path)
    create_lab_manual(repo2_path)
    print("All docx files generated successfully!")
